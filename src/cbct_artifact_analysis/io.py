"""Strict, headerless-workbook ingestion for the private master chart."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from cbct_artifact_analysis.study import EXPECTED_SITE_COUNTS, SCAN_METADATA

SITE_CODE_PATTERN = re.compile(r"^S\d+(?:LB|UB|LF|B|F)$", re.IGNORECASE)
SCAN_PATTERN = re.compile(r"Scan\s+([A-J])", re.IGNORECASE)
MEASUREMENT_COLUMNS = 72
BLOCK_SIZE = 24
BLOCK_COUNT = 3


class WorkbookValidationError(ValueError):
    """Raised when the private workbook no longer matches the audited structure."""


@dataclass(frozen=True)
class WorkbookProfile:
    filename: str
    sheet_name: str
    worksheet_rows: int
    worksheet_columns: int
    data_rows: int
    note_rows: int
    merged_ranges: int
    formulas: int
    comments: int
    tables: int
    charts: int
    hidden_rows: int
    hidden_columns: int


@dataclass(frozen=True)
class WorkbookData:
    long: pd.DataFrame
    sites: pd.DataFrame
    profile: WorkbookProfile


def _scan_id_from_label(label: object) -> str | None:
    if not isinstance(label, str):
        return None
    match = SCAN_PATTERN.search(label)
    return match.group(1).upper() if match else None


def load_master_workbook(path: str | Path) -> WorkbookData:
    """Load and validate the master chart without writing row-level data to disk."""

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(
            f"Private workbook not found: {workbook_path}. "
            "See data/README.md for the expected location."
        )

    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    if len(workbook.worksheets) != 1:
        raise WorkbookValidationError(f"Expected one worksheet, found {len(workbook.worksheets)}.")
    worksheet = workbook.worksheets[0]
    rows = list(worksheet.iter_rows(values_only=True))
    raw = pd.DataFrame(rows)

    if raw.shape[1] < 2 + MEASUREMENT_COLUMNS:
        raise WorkbookValidationError(f"Expected at least 74 columns, found {raw.shape[1]}.")

    site_mask = raw.iloc[:, 1].map(
        lambda value: isinstance(value, str) and bool(SITE_CODE_PATTERN.fullmatch(value.strip()))
    )
    site_rows = raw.loc[site_mask, : 1 + MEASUREMENT_COLUMNS].copy()
    if site_rows.empty:
        raise WorkbookValidationError("No site-code rows were found in column B.")

    # Merged cells make the scan label blank on continuation rows. Forward filling is safe
    # only inside the contiguous data region identified by site codes.
    site_rows.iloc[:, 0] = raw.iloc[: site_rows.index.max() + 1, 0].ffill().loc[site_rows.index]
    scan_ids = site_rows.iloc[:, 0].map(_scan_id_from_label)
    if scan_ids.isna().any():
        bad_rows = (scan_ids[scan_ids.isna()].index + 1).tolist()
        raise WorkbookValidationError(f"Could not infer scan IDs on worksheet rows {bad_rows}.")

    actual_scan_ids = set(scan_ids)
    expected_scan_ids = set(SCAN_METADATA)
    if actual_scan_ids != expected_scan_ids:
        raise WorkbookValidationError(
            f"Expected scans {sorted(expected_scan_ids)}, found {sorted(actual_scan_ids)}."
        )

    site_counts = scan_ids.value_counts().sort_index().to_dict()
    if site_counts != EXPECTED_SITE_COUNTS:
        raise WorkbookValidationError(
            f"Site counts changed. Expected {EXPECTED_SITE_COUNTS}, found {site_counts}."
        )

    numeric = site_rows.iloc[:, 2 : 2 + MEASUREMENT_COLUMNS].apply(pd.to_numeric, errors="coerce")
    if numeric.isna().any().any():
        positions = np.argwhere(numeric.isna().to_numpy())
        preview = [
            f"row {site_rows.index[row] + 1}, col {column + 3}" for row, column in positions[:8]
        ]
        raise WorkbookValidationError(
            "Non-numeric or missing measurement cells found at " + ", ".join(preview)
        )
    if not np.isfinite(numeric.to_numpy(dtype=float)).all():
        raise WorkbookValidationError("Infinite measurement values were found.")

    sites = pd.DataFrame(
        {
            "worksheet_row": site_rows.index + 1,
            "scan_id": scan_ids.to_numpy(),
            "site_code": site_rows.iloc[:, 1].astype(str).str.strip().to_numpy(),
            "source_label": site_rows.iloc[:, 0].astype(str).str.strip().to_numpy(),
        }
    )
    sites["site_number_within_scan"] = sites.groupby("scan_id").cumcount() + 1

    values = numeric.to_numpy(dtype=float).reshape(len(sites), BLOCK_COUNT, BLOCK_SIZE)
    long = pd.DataFrame(
        {
            "scan_id": np.repeat(sites["scan_id"].to_numpy(), MEASUREMENT_COLUMNS),
            "site_code": np.repeat(sites["site_code"].to_numpy(), MEASUREMENT_COLUMNS),
            "site_number_within_scan": np.repeat(
                sites["site_number_within_scan"].to_numpy(), MEASUREMENT_COLUMNS
            ),
            "block": np.tile(np.repeat(np.arange(1, BLOCK_COUNT + 1), BLOCK_SIZE), len(sites)),
            "measurement": np.tile(np.arange(1, BLOCK_SIZE + 1), len(sites) * BLOCK_COUNT),
            "value": values.reshape(-1),
        }
    )
    long["condition"] = long["scan_id"].map(
        {scan_id: metadata.condition for scan_id, metadata in SCAN_METADATA.items()}
    )
    long["alginate"] = long["scan_id"].map(
        {scan_id: metadata.alginate for scan_id, metadata in SCAN_METADATA.items()}
    )
    long["implant_count"] = long["scan_id"].map(
        {scan_id: metadata.implant_count for scan_id, metadata in SCAN_METADATA.items()}
    )

    data_end = int(site_rows.index.max()) + 1
    note_rows = sum(any(value is not None for value in row) for row in rows[data_end:])
    profile = WorkbookProfile(
        filename=workbook_path.name,
        sheet_name=worksheet.title,
        worksheet_rows=worksheet.max_row,
        worksheet_columns=worksheet.max_column,
        data_rows=len(sites),
        note_rows=note_rows,
        merged_ranges=len(worksheet.merged_cells.ranges),
        formulas=sum(cell.data_type == "f" for row in worksheet for cell in row),
        comments=sum(cell.comment is not None for row in worksheet for cell in row),
        tables=len(worksheet.tables),
        charts=len(worksheet._charts),
        hidden_rows=sum(bool(dimension.hidden) for dimension in worksheet.row_dimensions.values()),
        hidden_columns=sum(
            bool(dimension.hidden) for dimension in worksheet.column_dimensions.values()
        ),
    )
    return WorkbookData(long=long, sites=sites, profile=profile)
