from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cbct_artifact_analysis.study import EXPECTED_SITE_COUNTS

SITE_CODES = {
    "A": ["S1LB"],
    "B": ["S2LB"],
    "C": ["S3LB"],
    "D": ["S4LB"],
    "E": ["S5LB", "S5UB"],
    "F": ["S10LB", "S10UB"],
    "G": ["S6LB", "S6LF"],
    "H": ["S7B", "S7F"],
    "I": ["S8LB", "S8LF", "S8UB"],
    "J": ["S9LB", "S9LF", "S9UB"],
}


@pytest.fixture
def synthetic_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"

    row = 1
    for scan_id, site_count in EXPECTED_SITE_COUNTS.items():
        first_row = row
        for site_index in range(site_count):
            if site_index == 0:
                sheet.cell(row, 1, f"Scan {scan_id} - synthetic condition")
            sheet.cell(row, 2, SITE_CODES[scan_id][site_index])
            for offset in range(72):
                sheet.cell(row, 3 + offset, float(row * 1000 + offset))
            row += 1
        if site_count > 1:
            sheet.merge_cells(start_row=first_row, start_column=1, end_row=row - 1, end_column=1)
    sheet.cell(22, 1, "Free-text analysis note")
    workbook.save(path)
    return path
