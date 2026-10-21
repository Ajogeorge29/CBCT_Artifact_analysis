"""Create revised-framing copies while preserving the original manuscript files."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document

OBJECTIVE = (
    "Objective: To characterize and compare peri-implant CBCT gray-value distributions at "
    "standardized sampling locations around titanium implants in matched scan configurations "
    "acquired with and without alginate, and to assess consistency across three trials and tray "
    "status."
)

HYPOTHESIS = (
    "Study hypothesis: The null hypothesis is that matched alginate and no-alginate scans have "
    "the same peri-implant gray-value distribution. The research hypothesis is that the "
    "distributions differ. Direction and magnitude will be estimated rather than presupposed; "
    "gray-value differences will not be interpreted as direct proof of artifact reduction or "
    "worsening."
)

RESULTS = (
    "Results: The spreadsheet audit confirmed 18 site rows with 72 measurements per site, "
    "organized as three trials of 24 measurements. Nine of ten manuscript scan means reproduced "
    "from trial 1 at two decimals; Scan F was 632.2047 in the spreadsheet (632.20) versus 632.21 "
    "in the manuscripts. Across the matched B–C, E–F, G–H, and I–J configurations, alginate scans "
    "had higher descriptive mean peri-implant gray values in each trial. Because repeated points "
    "within a scan are correlated, these summaries are exploratory and do not independently "
    "establish artifact magnitude."
)

CONCLUSION = (
    "Conclusion: In this ex-vivo dataset, matched scans acquired with alginate showed higher "
    "descriptive peri-implant gray values across three trials. CBCT gray values are "
    "protocol-dependent surrogate imaging outcomes, and repeated sampling locations within a scan "
    "are not independent experimental replicates. These findings should not be interpreted as "
    "demonstrating that alginate reduces or worsens metal artifact. Independent acquisition-level "
    "replication and a validated direct artifact outcome are needed for confirmatory inference."
)


def _replace_first(document: Document, prefix: str, replacement: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            paragraph.text = replacement
            return
    raise ValueError(f"Could not find paragraph beginning with {prefix!r}")


def revise_document(source: Path, output: Path, full_manuscript: bool) -> None:
    document = Document(source)
    _replace_first(document, "Objective:", OBJECTIVE)
    _replace_first(document, "Results:", RESULTS)
    _replace_first(document, "Conclusion:", CONCLUSION)
    if full_manuscript:
        _replace_first(document, "The hypothesis of this study is", HYPOTHESIS)

    combined_text = "\n".join(paragraph.text for paragraph in document.paragraphs).lower()
    unsupported = [
        term
        for term in ("optical scanning", "dimensional shrinkage", "dimensional accuracy")
        if term in combined_text
    ]
    if unsupported:
        raise ValueError(f"Unsupported framing remains in {source.name}: {unsupported}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=Path("docs/source"))
    parser.add_argument("--output-dir", type=Path, default=Path("docs/revised"))
    args = parser.parse_args()

    jobs = [
        (
            args.source_dir / "Alginate April 2026.docx",
            args.output_dir / "Alginate April 2026 - revised framing.docx",
            True,
        ),
        (
            args.source_dir / "Alginate AADOCR Abstract.docx",
            args.output_dir / "Alginate AADOCR Abstract - revised framing.docx",
            False,
        ),
    ]
    for source, output, full_manuscript in jobs:
        revise_document(source, output, full_manuscript)
        print(f"Wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
